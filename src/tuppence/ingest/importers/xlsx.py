"""Excel workbooks: the first sheet with data, read cell by cell (runs in the sandbox)."""

from __future__ import annotations

import zipfile
from datetime import date, datetime

from tuppence.core.errors import UserFacing
from tuppence.ingest.refusals import Refusal

MAX_ROWS = 20_000
MAX_COLUMNS = 200
MAX_ENTRIES = 1_000
MAX_UNPACKED_BYTES = 200 * 1024 * 1024  # a workbook is a zip: refuse a zip bomb


class WorkbookRejected(UserFacing, ValueError):
    """The workbook can't be opened safely."""


class NotAWorkbook(Refusal, WorkbookRejected):
    message = "That doesn't look like an Excel workbook."


class WorkbookTooLarge(Refusal, WorkbookRejected):
    message = "That workbook is too large to read safely."


class WorkbookTooManyRows(Refusal, WorkbookRejected):
    message = "That workbook has too many rows to read safely."


def _check_zip(path: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
    except zipfile.BadZipFile:
        raise NotAWorkbook from None
    if len(infos) > MAX_ENTRIES or sum(i.file_size for i in infos) > MAX_UNPACKED_BYTES:
        raise WorkbookTooLarge


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return f"{value:.2f}" if round(value, 2) == value else repr(value)
    return " ".join(str(value).split())


def xlsx_records(path: str) -> list[tuple[int, str, list[str]]]:
    """(row number, line text, cells) for each row of the first sheet that has data."""
    import openpyxl

    _check_zip(path)
    try:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except (KeyError, ValueError, OSError, SyntaxError):  # SyntaxError: XML ParseError
        raise NotAWorkbook from None
    try:
        for sheet in book.worksheets:
            out: list[tuple[int, str, list[str]]] = []
            for n, row in enumerate(
                sheet.iter_rows(max_col=MAX_COLUMNS, values_only=True), start=1
            ):
                cells = [cell_text(v) for v in row]
                if n > MAX_ROWS:
                    if any(cells):
                        raise WorkbookTooManyRows
                    break
                while cells and not cells[-1]:
                    cells.pop()
                out.append((n, ", ".join(cells), cells))
            if any(cells for _, _, cells in out):
                return out
        return []
    finally:
        book.close()
